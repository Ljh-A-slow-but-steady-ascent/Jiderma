"""Generate the editable Jiderma project presentation."""

from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "artifacts" / "presentation"
OUT_DIR.mkdir(parents=True, exist_ok=True)
OUTPUT = OUT_DIR / "霁肤_Jiderma_项目汇报.pptx"
TRAINING_FIG = ROOT / "artifacts" / "training" / "ham7_linear_head" / "training_evaluation.png"
SAMPLE_IMAGE = ROOT / "third_party" / "DermFM-Zero" / "examples" / "PAT_8_15_820.png"

W, H = Inches(13.333), Inches(7.5)
FONT_CN = "Noto Sans CJK SC"
FONT_EN = "Aptos"

INK = RGBColor(24, 43, 56)
INK_2 = RGBColor(61, 81, 92)
MUTED = RGBColor(111, 130, 139)
BG = RGBColor(247, 250, 249)
WHITE = RGBColor(255, 255, 255)
LINE = RGBColor(216, 228, 226)
TEAL = RGBColor(38, 126, 126)
TEAL_D = RGBColor(24, 91, 94)
TEAL_L = RGBColor(226, 241, 239)
SKY = RGBColor(83, 151, 191)
SKY_L = RGBColor(229, 240, 247)
GOLD = RGBColor(219, 159, 67)
GOLD_L = RGBColor(249, 240, 220)
CORAL = RGBColor(211, 102, 84)
CORAL_L = RGBColor(250, 234, 230)
GREEN = RGBColor(72, 144, 108)
GREEN_L = RGBColor(232, 243, 236)


prs = Presentation()
prs.slide_width = W
prs.slide_height = H


def bg(slide, color=BG):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def rect(slide, x, y, w, h, fill=WHITE, line=LINE, radius=True, width=1):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(width)
    return shp


def line(slide, x1, y1, x2, y2, color=LINE, width=1.2, dash=False):
    shp = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    shp.line.color.rgb = color; shp.line.width = Pt(width)
    if dash:
        shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return shp


def text(slide, x, y, w, h, value, size=16, color=INK, bold=False,
         align=PP_ALIGN.LEFT, font=FONT_CN, valign=MSO_ANCHOR.TOP,
         margin=0.03, rotation=0):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.rotation = rotation
    tf = box.text_frame
    tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = value; p.alignment = align
    p.font.name = font; p.font.size = Pt(size); p.font.bold = bold; p.font.color.rgb = color
    return box


def rich(slide, x, y, w, h, runs, size=16, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(.03); tf.margin_top = tf.margin_bottom = Inches(.03)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]; p.alignment = align
    for content, opts in runs:
        r = p.add_run(); r.text = content
        r.font.name = opts.get("font", FONT_CN)
        r.font.size = Pt(opts.get("size", size))
        r.font.bold = opts.get("bold", False)
        r.font.color.rgb = opts.get("color", INK)
    return box


def circle(slide, x, y, d, fill=TEAL, line_color=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line_color is None: shp.line.fill.background()
    else: shp.line.color.rgb = line_color
    return shp


def brand(slide, x=.58, y=.28, dark=False):
    color = WHITE if dark else INK
    circle(slide, x, y, .34, TEAL if not dark else RGBColor(111, 203, 191))
    text(slide, x, y+.01, .34, .29, "霁", 14, WHITE, True, PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    rich(slide, x+.45, y-.01, 2.4, .36, [
        ("霁肤", {"size": 16, "bold": True, "color": color}),
        ("  Jiderma", {"size": 10.5, "bold": True, "color": TEAL if not dark else RGBColor(150, 225, 216), "font": FONT_EN}),
    ], valign=MSO_ANCHOR.MIDDLE)


def footer(slide, n, source=None, dark=False):
    color = RGBColor(179, 197, 199) if dark else MUTED
    text(slide, .62, 7.14, 10.9, .18,
         source or "辅助筛查与风险分级，不替代医生诊断", 7.5, color)
    text(slide, 12.14, 7.09, .55, .22, f"{n:02d}", 8.5, color, True, PP_ALIGN.RIGHT, FONT_EN)


def title(slide, kicker, heading, sub=None, n=None, source=None):
    brand(slide)
    text(slide, .62, .86, 2.8, .26, kicker.upper(), 9, TEAL, True, font=FONT_EN)
    text(slide, .62, 1.10, 12.0, .58, heading, 26, INK, True)
    if sub:
        text(slide, .64, 1.70, 11.7, .40, sub, 11.5, MUTED)
    if n is not None: footer(slide, n, source)


def pill(slide, x, y, w, label, fill=TEAL_L, color=TEAL_D, size=9):
    rect(slide, x, y, w, .31, fill, None, True)
    text(slide, x, y+.015, w, .25, label, size, color, True, PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)


def number_chip(slide, x, y, num, label, fill=TEAL):
    circle(slide, x, y, .34, fill)
    text(slide, x, y+.015, .34, .27, str(num), 10, WHITE, True, PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    text(slide, x+.45, y-.005, 2.1, .36, label, 11, INK, True, valign=MSO_ANCHOR.MIDDLE)


def arrow(slide, x1, y1, x2, y2, color=TEAL, width=1.8):
    shp = line(slide, x1, y1, x2, y2, color, width)
    shp.line.end_arrowhead = True
    return shp


# 1 — Cover
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s, RGBColor(19, 53, 62))
rect(s, 8.62, -.55, 5.2, 8.2, RGBColor(29, 72, 77), None, False)
circle(s, 9.15, .55, 3.55, RGBColor(43, 101, 101))
circle(s, 9.72, 1.12, 2.42, RGBColor(232, 190, 106))
circle(s, 10.28, 1.68, 1.30, RGBColor(249, 244, 218))
for i, yy in enumerate([4.65, 5.03, 5.41, 5.79]):
    line(s, 8.95 + i*.14, yy, 12.70, yy-.18, RGBColor(78, 139, 136), 1.0)
brand(s, .74, .52, True)
text(s, .78, 1.80, 7.15, .78, "霁肤 · Jiderma", 38, WHITE, True)
text(s, .80, 2.66, 7.0, .48, "面向基层医疗的皮肤疾病辅助决策 Agent", 20, RGBColor(198, 225, 220), True)
line(s, .80, 3.38, 2.02, 3.38, RGBColor(111, 203, 191), 3)
text(s, .80, 3.62, 6.88, 1.12,
     "从图像识别走向多源证据融合，\n让初筛更清晰，让转诊更有依据。", 22, WHITE, False)
pill(s, .80, 5.42, 1.34, "视觉基础模型", RGBColor(40, 83, 88), RGBColor(192, 232, 226))
pill(s, 2.27, 5.42, 1.18, "本地 RAG", RGBColor(40, 83, 88), RGBColor(192, 232, 226))
pill(s, 3.58, 5.42, 1.30, "受约束 Agent", RGBColor(40, 83, 88), RGBColor(192, 232, 226))
pill(s, 5.01, 5.42, 1.18, "MCP 工具", RGBColor(40, 83, 88), RGBColor(192, 232, 226))
text(s, .80, 6.56, 6.0, .28, "阶段成果与未来发展汇报  ·  2026", 10, RGBColor(164, 199, 196))
footer(s, 1, "霁＝雨后初晴，寓意拨云见日、恢复健康", True)


# 2 — Motivation and policy
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
title(s, "WHY NOW", "基层业务需求与政策窗口正在汇合",
      "皮肤问题高频、专科资源分布不均；基层真正需要的是“能识别、会核验、可转诊”的连续支持。", 2,
      "政策来源：国家卫健委《“十四五”全民健康信息化规划》、紧密型县域医共体建设指导意见、《人工智能+医疗卫生》实施意见")
rect(s, .62, 2.28, 4.18, 4.36, WHITE, LINE)
text(s, .92, 2.58, 3.55, .40, "一线场景的三个断点", 17, INK, True)
items = [
    ("01", "看不准", "皮肤表现相似，单张图片难以覆盖病程与风险信号", CORAL_L, CORAL),
    ("02", "问不全", "基层问诊时间有限，症状与变化信息容易遗漏", GOLD_L, GOLD),
    ("03", "接不上", "分类结果、指南依据与转诊行动常处于不同系统", SKY_L, SKY),
]
for i, (num, hd, body, fc, ac) in enumerate(items):
    yy = 3.24 + i*1.02
    circle(s, .92, yy, .46, fc)
    text(s, .92, yy+.04, .46, .30, num, 9, ac, True, PP_ALIGN.CENTER, FONT_EN, MSO_ANCHOR.MIDDLE)
    text(s, 1.56, yy-.02, .85, .28, hd, 13, INK, True)
    text(s, 2.40, yy-.02, 2.02, .55, body, 9.5, MUTED)

text(s, 5.22, 2.34, 7.45, .36, "政策节奏：从数字化基层，到智能辅助广泛应用", 16, INK, True)
line(s, 5.48, 3.18, 12.25, 3.18, RGBColor(181, 208, 206), 2)
policy = [
    (5.34, "2022", "基层信息能力", "推动基层智能辅助诊断、线上复诊与转诊协同"),
    (7.48, "2023", "县域医共体", "县级强、乡级活、村级稳、上下联、信息通"),
    (9.62, "2025", "AI+医疗卫生", "专病模型与智能体、基层辅助诊疗进入规模化窗口"),
]
for x, yr, hd, bd in policy:
    circle(s, x, 2.96, .44, TEAL)
    text(s, x-.03, 3.52, 1.0, .30, yr, 11, TEAL_D, True, font=FONT_EN)
    text(s, x-.03, 3.90, 1.78, .34, hd, 12, INK, True)
    text(s, x-.03, 4.33, 1.92, .84, bd, 9.4, MUTED)
rect(s, 5.24, 5.52, 7.40, .94, TEAL_D, None)
rich(s, 5.55, 5.74, 6.80, .50, [
    ("霁肤的落点：", {"size": 14, "bold": True, "color": WHITE}),
    ("把专科能力封装为基层可用的筛查、证据与转诊工作流。", {"size": 14, "color": RGBColor(218, 240, 236)}),
], valign=MSO_ANCHOR.MIDDLE)


# 3 — Product proposition
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
title(s, "PRODUCT", "不是“给一个标签”，而是完成一次辅助决策",
      "输入图像与主诉，系统组织视觉、症状、规则和指南证据，输出风险解释与下一步行动。", 3)
text(s, .65, 2.28, 4.20, .30, "传统图像分类", 13, MUTED, True)
rect(s, .62, 2.70, 4.15, 2.98, RGBColor(242, 245, 245), LINE)
text(s, 1.00, 3.05, 3.4, .44, "图片  →  Top-1 类别", 22, INK_2, True, PP_ALIGN.CENTER)
text(s, 1.02, 3.77, 3.35, 1.12,
     "缺少病程与症状语境\n缺少风险冲突判断\n缺少可追溯行动依据", 13, MUTED, False, PP_ALIGN.CENTER)
text(s, 5.19, 2.28, 4.20, .30, "霁肤 · Jiderma", 13, TEAL, True)
rect(s, 5.16, 2.70, 7.52, 2.98, WHITE, TEAL, True, 1.6)
steps = [
    (5.48, "01", "图像", "候选识别"), (6.85, "02", "症状", "结构化核验"),
    (8.22, "03", "冲突", "风险升级"), (9.59, "04", "证据", "指南检索"),
    (10.96, "05", "行动", "照护建议"),
]
for i, (x, no, a, b) in enumerate(steps):
    circle(s, x, 3.30, .70, TEAL_L)
    text(s, x, 3.44, .70, .28, no, 10, TEAL_D, True, PP_ALIGN.CENTER, FONT_EN, MSO_ANCHOR.MIDDLE)
    text(s, x-.12, 4.22, .94, .28, a, 12, INK, True, PP_ALIGN.CENTER)
    text(s, x-.20, 4.58, 1.10, .30, b, 9.5, MUTED, False, PP_ALIGN.CENTER)
    if i < len(steps)-1: arrow(s, x+.76, 3.66, x+1.21, 3.66, RGBColor(150, 190, 186), 1.2)
rect(s, 2.44, 6.06, 8.46, .58, TEAL_L, None)
text(s, 2.62, 6.19, 8.10, .28, "核心价值：从“模型输出”跨越到“基层可执行的辅助决策”", 14, TEAL_D, True, PP_ALIGN.CENTER)


# 4 — Experiment preparation
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
title(s, "VALIDATION", "轻量任务头验证：低成本激活专业视觉表征",
      "冻结视觉 Backbone，仅训练 Linear Head；以较小训练成本验证下游任务迁移能力。", 4)
rect(s, .62, 2.30, 5.10, 3.98, WHITE, LINE)
text(s, .93, 2.62, 4.45, .32, "实验路径", 16, INK, True)
number_chip(s, .93, 3.20, 1, "专业视觉表征", SKY)
number_chip(s, .93, 3.90, 2, "冻结 Backbone", TEAL)
number_chip(s, .93, 4.60, 3, "训练 Linear Head", GOLD)
number_chip(s, .93, 5.30, 4, "独立测试评估", GREEN)
text(s, 3.44, 3.15, 1.80, .50, "768 维特征", 18, SKY, True, PP_ALIGN.RIGHT, FONT_EN)
text(s, 3.44, 4.02, 1.80, .50, "5,383", 25, TEAL_D, True, PP_ALIGN.RIGHT, FONT_EN)
text(s, 3.44, 4.52, 1.80, .26, "可训练参数", 9, MUTED, False, PP_ALIGN.RIGHT)
text(s, 6.12, 2.44, 2.30, .32, "数据划分", 16, INK, True)
splits = [("TRAIN", "8,012", 0.801, TEAL), ("VALIDATION", "500", .050, GOLD), ("TEST", "1,503", .150, SKY)]
for i, (lab, val, frac, col) in enumerate(splits):
    yy = 3.18 + i*.84
    text(s, 6.14, yy, 1.30, .26, lab, 9, MUTED, True, font=FONT_EN)
    rect(s, 7.46, yy+.01, 3.55, .20, RGBColor(231, 237, 236), None, True)
    rect(s, 7.46, yy+.01, 3.55*frac, .20, col, None, True)
    text(s, 11.24, yy-.06, .82, .30, val, 13, INK, True, PP_ALIGN.RIGHT, FONT_EN)
rect(s, 6.12, 5.85, 6.15, .52, SKY_L, None)
text(s, 6.33, 5.97, 5.72, .27, "一套 Backbone，可继续扩展多病种与多任务头", 12, SKY, True, PP_ALIGN.CENTER)


# 5 — Model architecture
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
title(s, "MODEL ARCHITECTURE", "从皮肤图像到结构化视觉证据",
      "专业视觉基础模型负责表征，轻量任务头完成场景适配；最终候选交由决策平面继续核验。", 5,
      "技术依据：DermFM-Zero-Open 本地代码与当前 HAM10000 线性任务头实现")

# Main inference pipeline
text(s, .66, 2.20, 2.10, .28, "INFERENCE PATH", 9, TEAL, True, font=FONT_EN)
rect(s, .64, 2.55, 1.72, 2.38, WHITE, LINE)
if SAMPLE_IMAGE.exists():
    s.shapes.add_picture(str(SAMPLE_IMAGE), Inches(.83), Inches(2.78), width=Inches(1.34), height=Inches(1.34))
text(s, .83, 4.25, 1.34, .27, "病灶图像", 12, INK, True, PP_ALIGN.CENTER)
text(s, .79, 4.56, 1.42, .20, "Resize → Crop → Normalize", 7.5, MUTED, False, PP_ALIGN.CENTER, FONT_EN)
arrow(s, 2.39, 3.68, 2.77, 3.68, TEAL, 1.7)

rect(s, 2.82, 2.55, 3.20, 2.38, TEAL_D, None)
pill(s, 3.08, 2.81, 1.06, "FROZEN", RGBColor(52, 112, 112), RGBColor(210, 239, 234), 8)
text(s, 3.08, 3.32, 2.69, .35, "DermFM-Zero-Open", 16, WHITE, True, font=FONT_EN)
text(s, 3.08, 3.76, 2.70, .27, "Vision Transformer Encoder", 10, RGBColor(194, 225, 221), True, font=FONT_EN)
# Transformer token motif
for row in range(2):
    for col in range(6):
        rect(s, 3.10+col*.39, 4.23+row*.24, .27, .14,
             RGBColor(72+row*12, 139+col*3, 137+col*3), None, False)
arrow(s, 6.07, 3.68, 6.43, 3.68, TEAL, 1.7)

rect(s, 6.48, 2.55, 1.60, 2.38, SKY_L, SKY)
text(s, 6.68, 2.86, 1.20, .27, "视觉表征", 12, SKY, True, PP_ALIGN.CENTER)
text(s, 6.67, 3.43, 1.22, .56, "768-D", 25, INK, True, PP_ALIGN.CENTER, FONT_EN)
text(s, 6.66, 4.08, 1.24, .48, "L2-normalized\nembedding", 8.5, MUTED, False, PP_ALIGN.CENTER, FONT_EN)
arrow(s, 8.13, 3.68, 8.49, 3.68, TEAL, 1.7)

rect(s, 8.54, 2.55, 2.02, 2.38, GOLD_L, GOLD)
pill(s, 8.78, 2.81, 1.50, "TRAINABLE", RGBColor(245, 226, 186), RGBColor(174, 115, 31), 8)
text(s, 8.79, 3.34, 1.52, .29, "Linear Head", 15, INK, True, PP_ALIGN.CENTER, FONT_EN)
text(s, 8.76, 3.80, 1.56, .34, "W · z + b", 17, GOLD, True, PP_ALIGN.CENTER, FONT_EN)
text(s, 8.72, 4.27, 1.64, .27, "768 → 7  ·  5,383 params", 8.2, MUTED, False, PP_ALIGN.CENTER, FONT_EN)
arrow(s, 10.61, 3.68, 10.96, 3.68, TEAL, 1.7)

rect(s, 11.01, 2.55, 1.66, 2.38, WHITE, LINE)
text(s, 11.22, 2.84, 1.22, .28, "候选输出", 12, INK, True, PP_ALIGN.CENTER)
bars = [("Top-1", .94, TEAL), ("Top-2", .52, SKY), ("Top-3", .28, GOLD)]
for i, (lab, frac, col) in enumerate(bars):
    yy = 3.37 + i*.43
    text(s, 11.17, yy-.01, .45, .17, lab, 7.5, MUTED, True, font=FONT_EN)
    rect(s, 11.65, yy, .76, .12, RGBColor(232, 238, 237), None, True)
    rect(s, 11.65, yy, .76*frac, .12, col, None, True)
text(s, 11.18, 4.65, 1.30, .17, "temperature + abstention", 7, MUTED, False, PP_ALIGN.CENTER, FONT_EN)

# Lower band: training rationale and interface boundary
rect(s, .64, 5.34, 7.43, 1.15, WHITE, LINE)
text(s, .90, 5.58, 1.25, .26, "训练策略", 12, TEAL_D, True)
rich(s, 2.02, 5.55, 5.75, .35, [
    ("冻结编码器", {"size": 11, "bold": True, "color": INK}),
    ("  +  HAM10000 监督训练  +  ", {"size": 11, "color": MUTED}),
    ("温度校准 / 拒识阈值", {"size": 11, "bold": True, "color": INK}),
])
text(s, 2.02, 5.98, 5.63, .26, "预训练表征：300万+皮肤图像表征学习，100万图文对齐样本", 9.2, SKY, True)

rect(s, 8.42, 5.34, 4.25, 1.15, TEAL_L, None)
text(s, 8.69, 5.55, 1.18, .27, "接口边界", 12, TEAL_D, True)
text(s, 9.70, 5.50, 2.68, .56, "仅输出结构化视觉证据\n风险升级由症状与本地规则共同决定", 9.8, INK, True)


# 6 — Results
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
title(s, "RESULTS", "HAM10000 下游任务已完成独立测试",
      "实验覆盖训练日志、模型选择、测试指标、混淆矩阵与类别级表现。", 6)
metrics = [("89.29%", "Accuracy", TEAL), ("79.34%", "Balanced Accuracy", SKY), ("80.98%", "Macro-F1", GOLD)]
for i, (val, lab, col) in enumerate(metrics):
    x = .64 + i*2.18
    rect(s, x, 2.30, 1.94, 1.03, WHITE, LINE)
    text(s, x+.12, 2.46, 1.70, .40, val, 22, col, True, PP_ALIGN.CENTER, FONT_EN)
    text(s, x+.10, 2.94, 1.74, .22, lab, 8.5, MUTED, True, PP_ALIGN.CENTER, FONT_EN)
if TRAINING_FIG.exists():
    s.shapes.add_picture(str(TRAINING_FIG), Inches(6.88), Inches(2.23), width=Inches(5.80), height=Inches(4.27))
rect(s, .64, 3.64, 5.84, 2.86, TEAL_D, None)
text(s, .98, 3.98, 5.20, .36, "实验结论", 17, WHITE, True)
text(s, .98, 4.51, 5.08, 1.22,
     "专业视觉表征能够有效迁移到固定皮肤镜分类任务；\n轻量任务头使新任务适配更快、更节省算力，也更便于工程迭代。", 13, RGBColor(220, 240, 236))
pill(s, .98, 5.88, 1.33, "冻结式迁移", RGBColor(50, 109, 108), RGBColor(224, 243, 239))
pill(s, 2.46, 5.88, 1.33, "低训练成本", RGBColor(50, 109, 108), RGBColor(224, 243, 239))
pill(s, 3.94, 5.88, 1.63, "可扩展任务头", RGBColor(50, 109, 108), RGBColor(224, 243, 239))


# 7 — Workflow
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
title(s, "WORKFLOW", "一次完整辅助决策，如何发生？",
      "视觉提供候选，症状补全语境，规则守住风险边界，Agent 组织证据与表达。", 7)
labels = [
    ("01", "上传图像", "病灶图像 + 基本主诉", SKY_L, SKY),
    ("02", "视觉候选", "疾病候选与相对分数", TEAL_L, TEAL),
    ("03", "症状核验", "病程、变化与红旗信号", GOLD_L, GOLD),
    ("04", "风险冲突", "多源证据交叉判断", CORAL_L, CORAL),
    ("05", "指南检索", "动态 Query + 本地 RAG", GREEN_L, GREEN),
    ("06", "建议生成", "解释、分级与照护路径", SKY_L, SKY),
]
for i, (no, hd, bd, fc, ac) in enumerate(labels):
    x = .58 + i*2.10
    rect(s, x, 2.62, 1.78, 2.14, WHITE, LINE)
    circle(s, x+.16, 2.84, .43, fc)
    text(s, x+.16, 2.91, .43, .26, no, 9, ac, True, PP_ALIGN.CENTER, FONT_EN, MSO_ANCHOR.MIDDLE)
    text(s, x+.16, 3.50, 1.45, .35, hd, 13, INK, True)
    text(s, x+.16, 3.98, 1.43, .56, bd, 9.3, MUTED)
    if i < len(labels)-1: arrow(s, x+1.82, 3.67, x+2.05, 3.67, RGBColor(155, 185, 183), 1.3)
rect(s, 1.20, 5.34, 10.94, 1.08, WHITE, LINE)
text(s, 1.48, 5.58, 2.05, .28, "典型冲突场景", 12, CORAL, True)
rich(s, 3.24, 5.54, 8.52, .42, [
    ("视觉候选偏向低风险  ", {"size": 12, "color": INK}),
    ("但", {"size": 12, "bold": True, "color": MUTED}),
    (" 近期快速增大 / 出血  ", {"size": 12, "bold": True, "color": CORAL}),
    ("→ 触发风险升级与定向指南检索", {"size": 12, "bold": True, "color": TEAL_D}),
])


# 8 — Architecture
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
title(s, "ARCHITECTURE", "本地优先、双平面协同的系统架构",
      "技术点到为止：Foundation Model · Linear Probe · Hybrid RAG · ReAct · MCP · PrivFlow", 8)

# layer captions
layers = [
    (2.25, "交互接入", SKY_L, SKY, ["Streamlit Web", "图像上传", "悬浮 ChatBox"]),
    (3.24, "业务服务", TEAL_L, TEAL, ["FastAPI", "病例状态", "SQLite 审计"]),
    (4.23, "决策平面", GOLD_L, GOLD, ["视觉候选", "症状核验", "风险规则"]),
    (5.22, "Agent 平面", GREEN_L, GREEN, ["ReAct 规划", "MCP 工具", "本地医学 RAG"]),
]
for yy, cap, fc, ac, nodes in layers:
    rect(s, .62, yy, 12.06, .78, WHITE, LINE)
    rect(s, .62, yy, 1.38, .78, fc, None, False)
    text(s, .75, yy+.19, 1.10, .30, cap, 11, ac, True, PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    for j, node in enumerate(nodes):
        xx = 2.24 + j*2.82
        pill(s, xx, yy+.22, 2.34, node, fc, ac, 10)
    if yy < 5.2: arrow(s, 6.60, yy+.80, 6.60, yy+.97, RGBColor(144, 177, 175), 1.2)

rect(s, 10.52, 3.47, 1.72, 2.18, RGBColor(31, 73, 79), None)
text(s, 10.68, 3.73, 1.40, .30, "安全治理侧车", 11, WHITE, True, PP_ALIGN.CENTER)
text(s, 10.66, 4.20, 1.43, 1.08,
     "隐私脱敏\n出境检查\n来源追溯\n服务降级", 10.5, RGBColor(207, 233, 229), False, PP_ALIGN.CENTER)
text(s, .68, 6.36, 11.92, .30,
     "关键原则：LLM 负责检索规划与解释；风险分级由本地可审计规则约束。", 12.5, TEAL_D, True, PP_ALIGN.CENTER)


# 9 — Differentiation
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
title(s, "DIFFERENTIATION", "差异化不在模型堆叠，而在决策链闭环",
      "感知、证据、风险与行动被组织为同一条可追溯工作流。", 9)
diffs = [
    ("01", "专业视觉底座", "从通用视觉理解转向皮肤场景表征，任务头可按业务扩展。", SKY_L, SKY),
    ("02", "多源证据融合", "图像、症状、病程、风险规则和医学指南共同参与判断。", TEAL_L, TEAL),
    ("03", "动态可追溯 RAG", "Agent 围绕信息缺口检索，并保留机构、章节与来源。", GREEN_L, GREEN),
    ("04", "受约束 Agent", "关键风险不交给自由生成，输出具备边界与审计能力。", GOLD_L, GOLD),
    ("05", "轻量可落地", "本地优先、服务可降级，适配基层、社区与私有化部署。", CORAL_L, CORAL),
]
for i, (no, hd, bd, fc, ac) in enumerate(diffs):
    col, row = i % 3, i // 3
    x = .63 + col*4.18
    y = 2.40 + row*1.82
    w = 3.80 if i < 3 else 5.92
    if i >= 3: x = .63 + (i-3)*6.22
    rect(s, x, y, w, 1.48, WHITE, LINE)
    circle(s, x+.22, y+.24, .44, fc)
    text(s, x+.22, y+.31, .44, .26, no, 9, ac, True, PP_ALIGN.CENTER, FONT_EN, MSO_ANCHOR.MIDDLE)
    text(s, x+.84, y+.22, w-1.10, .30, hd, 14, INK, True)
    text(s, x+.84, y+.67, w-1.14, .57, bd, 9.7, MUTED)
rect(s, 3.20, 6.23, 6.94, .47, TEAL_D, None)
text(s, 3.42, 6.33, 6.50, .24, "识别 → 核验 → 检索 → 风险判断 → 行动建议", 12, WHITE, True, PP_ALIGN.CENTER)


# 10 — Video demo placeholder
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s, RGBColor(23, 51, 60))
brand(s, .58, .28, True)
text(s, .64, .92, 11.8, .58, "产品演示 · 从一张图到一条照护路径", 26, WHITE, True)
text(s, .66, 1.56, 11.7, .34, "预留视频位：建议插入 60–90 秒真实操作录屏", 11.5, RGBColor(173, 205, 203))
rect(s, .66, 2.22, 8.20, 4.32, RGBColor(13, 34, 42), RGBColor(72, 121, 122), True, 1.2)
circle(s, 4.06, 3.55, 1.42, RGBColor(55, 133, 130))
tri = s.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, Inches(4.58), Inches(3.93), Inches(.40), Inches(.48))
tri.rotation = 90; tri.fill.solid(); tri.fill.fore_color.rgb = WHITE; tri.line.fill.background()
text(s, 2.22, 5.42, 5.08, .32, "VIDEO PLACEHOLDER", 12, RGBColor(139, 176, 175), True, PP_ALIGN.CENTER, FONT_EN)
text(s, 9.28, 2.35, 3.30, .34, "演示脚本", 16, WHITE, True)
demo = ["上传病灶图像", "补充结构化症状", "查看候选与冲突判断", "追问风险与疾病知识", "获得证据化照护建议"]
for i, item in enumerate(demo, 1):
    yy = 3.00 + (i-1)*.67
    circle(s, 9.30, yy, .34, RGBColor(49, 104, 105))
    text(s, 9.30, yy+.02, .34, .26, str(i), 9, RGBColor(193, 230, 224), True, PP_ALIGN.CENTER, FONT_EN, MSO_ANCHOR.MIDDLE)
    text(s, 9.78, yy-.01, 2.58, .34, item, 11, RGBColor(222, 238, 236), True, valign=MSO_ANCHOR.MIDDLE)
footer(s, 10, "视频文件可在 PowerPoint 中替换该占位框后直接播放", True)


# 11 — Roadmap
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
title(s, "ROADMAP", "从可演示原型，走向基层皮肤智能服务平台",
      "现有能力已经形成技术闭环；未来投入将主要转化为病种覆盖、临床协作与规模化交付。", 11)
stages = [
    ("NOW", "原型验证", "视觉任务头\n结构化症状\n本地 RAG / Agent", TEAL, TEAL_L),
    ("NEXT", "能力扩展", "多病种任务头\n多视角病例理解\n病灶变化追踪", SKY, SKY_L),
    ("SCALE", "临床协作", "医生工作台\n转诊材料生成\n区域知识更新", GOLD, GOLD_L),
    ("VISION", "平台化落地", "端侧 / 私有化部署\n连续随访管理\n县域协同网络", GREEN, GREEN_L),
]
for i, (tag, hd, bd, ac, fc) in enumerate(stages):
    x = .64 + i*3.12
    rect(s, x, 2.48, 2.74, 3.26, WHITE, LINE)
    pill(s, x+.22, 2.76, .78, tag, fc, ac, 8.5)
    text(s, x+.22, 3.36, 2.24, .38, hd, 16, INK, True)
    text(s, x+.22, 4.08, 2.26, 1.12, bd, 11, MUTED)
    if i < 3: arrow(s, x+2.78, 4.12, x+3.05, 4.12, RGBColor(146, 183, 180), 1.5)
rect(s, 1.25, 6.18, 10.82, .66, TEAL_D, None)
text(s, 1.52, 6.32, 10.28, .33,
     "霁肤 · Jiderma  ——  让基层初筛看得更清，走得更稳。", 16, WHITE, True, PP_ALIGN.CENTER)


# Metadata and save
prs.core_properties.title = "霁肤 · Jiderma 项目汇报"
prs.core_properties.subject = "面向基层医疗的皮肤疾病辅助决策 Agent"
prs.core_properties.author = "Jiderma Project Team"
prs.core_properties.keywords = "Jiderma, 霁肤, 基层医疗, RAG, Agent, MCP"
prs.save(OUTPUT)
print(OUTPUT)
