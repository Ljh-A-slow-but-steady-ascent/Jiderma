"""DermaGuard 项目 PPT 生成脚本。

基于项目.md 和 docs/视觉模型与训练.md 生成一份完整的项目介绍 PPTX。
深色科技主题，16:9 宽屏。
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ---------------------------------------------------------------------------
# 颜色常量（与项目深色主题一致）
# ---------------------------------------------------------------------------
BG_DARK = RGBColor(0x05, 0x08, 0x10)
BG_CARD = RGBColor(0x0E, 0x1A, 0x28)
CYAN = RGBColor(0x00, 0xF5, 0xD4)
MAGENTA = RGBColor(0xFF, 0x2D, 0x95)
YELLOW = RGBColor(0xFF, 0xD2, 0x3F)
BLUE = RGBColor(0x4D, 0xD0, 0xFF)
GREEN = RGBColor(0x00, 0xFF, 0x88)
RED = RGBColor(0xFF, 0x40, 0x60)
WHITE = RGBColor(0xD4, 0xF5, 0xFF)
MUTED = RGBColor(0x6A, 0x8A, 0x9C)
AMBER = RGBColor(0xFF, 0xAE, 0x3F)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H


# ---------------------------------------------------------------------------
# 辅助函数
# ---------------------------------------------------------------------------
def _set_bg(slide, color=BG_DARK):
    """设置幻灯片背景色。"""
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def _add_textbox(slide, left, top, width, height, text, size=18,
                 color=WHITE, bold=False, align=PP_ALIGN.LEFT, font_name="Microsoft YaHei"):
    """添加文本框。"""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = align
    return txBox


def _add_multi_text(slide, left, top, width, height, lines, size=16,
                    color=WHITE, font_name="Microsoft YaHei", line_spacing=1.5):
    """添加多行文本。"""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, (text, opts) in enumerate(lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = text
        p.font.size = Pt(opts.get("size", size))
        p.font.color.rgb = opts.get("color", color)
        p.font.bold = opts.get("bold", False)
        p.font.name = font_name
        p.alignment = opts.get("align", PP_ALIGN.LEFT)
        p.space_after = Pt(opts.get("space_after", 6))
        p.line_spacing = line_spacing
    return txBox


def _add_card(slide, left, top, width, height, color=BG_CARD):
    """添加圆角矩形卡片背景。"""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = RGBColor(0x00, 0x35, 0x2A)
    shape.line.width = Pt(1)
    shape.shadow.inherit = False
    return shape


def _add_accent_bar(slide, left, top, width, height, color=CYAN):
    """添加装饰条。"""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, left, top, width, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def _add_title_bar(slide, title, subtitle=""):
    """添加页面标题栏。"""
    _add_accent_bar(slide, Inches(0.6), Inches(0.4), Inches(0.08), Inches(0.55), CYAN)
    _add_textbox(slide, Inches(0.85), Inches(0.35), Inches(10), Inches(0.7),
                 title, size=28, color=CYAN, bold=True)
    if subtitle:
        _add_textbox(slide, Inches(0.85), Inches(0.95), Inches(10), Inches(0.4),
                     subtitle, size=14, color=MUTED)


def _add_bullet_list(slide, left, top, width, height, items, size=16, color=WHITE):
    """添加项目符号列表。"""
    lines = []
    for item in items:
        if isinstance(item, tuple):
            text, opts = item
        else:
            text, opts = item, {}
        lines.append((f"  ▸  {text}", opts))
    return _add_multi_text(slide, left, top, width, height, lines, size=size, color=color)


def _add_table(slide, left, top, width, height, headers, rows, header_color=CYAN):
    """添加表格。"""
    num_rows = len(rows) + 1
    num_cols = len(headers)
    table_shape = slide.shapes.add_table(num_rows, num_cols, left, top, width, height)
    table = table_shape.table

    # 表头
    for j, header in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = header
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(14)
            p.font.bold = True
            p.font.color.rgb = header_color
            p.font.name = "Microsoft YaHei"
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(0x0A, 0x14, 0x20)

    # 数据行
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.cell(i + 1, j)
            cell.text = str(val)
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(13)
                p.font.color.rgb = WHITE
                p.font.name = "Microsoft YaHei"
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(0x08, 0x10, 0x18) if i % 2 == 0 else RGBColor(0x0A, 0x14, 0x20)

    return table_shape


# ===========================================================================
# Slide 1: 封面
# ===========================================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
_set_bg(slide)

# 装饰条
_add_accent_bar(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.06), CYAN)
_add_accent_bar(slide, Inches(0), Inches(7.44), Inches(13.333), Inches(0.06), MAGENTA)

# 标题
_add_textbox(slide, Inches(1.5), Inches(1.8), Inches(10), Inches(1.2),
             "DermaGuard", size=54, color=CYAN, bold=True, align=PP_ALIGN.CENTER)
_add_textbox(slide, Inches(1.5), Inches(3.0), Inches(10), Inches(0.8),
             "基层皮肤病灶智能筛查与分级转诊演示系统", size=24, color=WHITE, align=PP_ALIGN.CENTER)

# 标签
_add_textbox(slide, Inches(2.5), Inches(4.0), Inches(8), Inches(0.5),
             "Local-First · 隐私保护 · AI 辅助 · MCP 工具链 · PrivFlow 治理",
             size=14, color=MAGENTA, align=PP_ALIGN.CENTER)

# 底部信息
_add_textbox(slide, Inches(1.5), Inches(5.5), Inches(10), Inches(0.5),
             "本地决策平面 · 云端解释平面 · 结构化症状核验 · 本地 RAG",
             size=13, color=MUTED, align=PP_ALIGN.CENTER)
_add_textbox(slide, Inches(1.5), Inches(6.2), Inches(10), Inches(0.5),
             "2026-10-11 · 完成度约 92%", size=12, color=MUTED, align=PP_ALIGN.CENTER)

# ===========================================================================
# Slide 2: 项目定位
# ===========================================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
_set_bg(slide)
_add_title_bar(slide, "项目定位", "Project Positioning")

# 核心定义卡片
_add_card(slide, Inches(0.85), Inches(1.4), Inches(11.5), Inches(1.6))
_add_multi_text(slide, Inches(1.1), Inches(1.55), Inches(11), Inches(1.3), [
    ("DermaGuard 是一个面向基层医疗场景的皮肤病灶智能筛查与分级转诊演示系统。",
     {"size": 18, "color": WHITE, "bold": True}),
    ("整合本地视觉模型、结构化症状核验、本地医学知识检索、确定性风险规则、照护路径、MCP 工具、隐私出境治理和云端大语言模型解释为一条可运行流程。",
     {"size": 15, "color": MUTED}),
])

# 三个特性卡
card_w = Inches(3.6)
card_h = Inches(2.2)
gap = Inches(0.3)
start_x = Inches(0.85)
y = Inches(3.4)

cards = [
    ("可运行", "完整端到端流程\n从图片上传到分级转诊\n全链路可演示", CYAN),
    ("职责清晰", "决策平面安全可控\n交互平面只做解释\n云端不可修改本地结果", MAGENTA),
    ("隐私保护", "原始影像不出域\nPrivFlow 出境治理\n全链路审计可追溯", YELLOW),
]
for i, (title, desc, color) in enumerate(cards):
    x = start_x + i * (card_w + gap)
    _add_card(slide, x, y, card_w, card_h)
    _add_accent_bar(slide, x, y, Inches(0.08), card_h, color)
    _add_textbox(slide, x + Inches(0.2), y + Inches(0.2), card_w - Inches(0.3), Inches(0.5),
                 title, size=20, color=color, bold=True)
    _add_textbox(slide, x + Inches(0.2), y + Inches(0.8), card_w - Inches(0.3), Inches(1.2),
                 desc, size=14, color=WHITE)

_add_textbox(slide, Inches(0.85), Inches(6.0), Inches(11), Inches(0.5),
             "本系统仅用于辅助筛查和分级转诊，不构成医学诊断。", size=13, color=RED)

# ===========================================================================
# Slide 3: 双平面架构
# ===========================================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
_set_bg(slide)
_add_title_bar(slide, "双平面架构", "Dual-Plane Architecture")

# 决策平面
_add_card(slide, Inches(0.85), Inches(1.3), Inches(5.6), Inches(5.5), RGBColor(0x04, 0x14, 0x1E))
_add_accent_bar(slide, Inches(0.85), Inches(1.3), Inches(5.6), Inches(0.08), CYAN)
_add_textbox(slide, Inches(1.1), Inches(1.5), Inches(5), Inches(0.5),
             "决策平面 · Decision Plane", size=20, color=CYAN, bold=True)
_add_textbox(slide, Inches(1.1), Inches(2.0), Inches(5), Inches(0.3),
             "负责安全关键结果", size=12, color=MUTED)

decision_flow = [
    "患者上传图片",
    "本地 DermFM-Zero 冻结主干 + HAM 监督任务头",
    "视觉候选、图像质量和运行元数据",
    "本地确定性 Triage Policy",
    "结构化症状核验 + 本地 RAG",
    "本地 Context Risk Policy 二次校准",
    "Care Pathway 照护路径",
]
for i, step in enumerate(decision_flow):
    y_pos = Inches(2.5 + i * 0.55)
    _add_textbox(slide, Inches(1.3), y_pos, Inches(4.8), Inches(0.4),
                 f"  {i+1}. {step}", size=13, color=WHITE if i < len(decision_flow)-1 else CYAN)
    if i < len(decision_flow) - 1:
        _add_textbox(slide, Inches(1.5), y_pos + Inches(0.38), Inches(4), Inches(0.2),
                     "    ↓", size=11, color=MUTED, align=PP_ALIGN.LEFT)

_add_textbox(slide, Inches(1.1), Inches(6.3), Inches(5), Inches(0.4),
             "云端 LLM 不能修改视觉分数、风险等级或照护路径", size=11, color=RED)

# 交互平面
_add_card(slide, Inches(6.85), Inches(1.3), Inches(5.6), Inches(5.5), RGBColor(0x1A, 0x04, 0x14))
_add_accent_bar(slide, Inches(6.85), Inches(1.3), Inches(5.6), Inches(0.08), MAGENTA)
_add_textbox(slide, Inches(7.1), Inches(1.5), Inches(5), Inches(0.5),
             "交互平面 · Interaction Plane", size=20, color=MAGENTA, bold=True)
_add_textbox(slide, Inches(7.1), Inches(2.0), Inches(5), Inches(0.3),
             "负责解释、问答与导航", size=12, color=MUTED)

interact_flow = [
    "用户提问",
    "DermaGuardChatAgent",
    "ToolRegistry / MCP Tools",
    "PrivFlow EgressGuard 出境检查",
    "智谱 GLM 或本地 Mock Provider",
    "自然语言解释",
    "Tool Call + 审计写入 SQLite",
]
for i, step in enumerate(interact_flow):
    y_pos = Inches(2.5 + i * 0.55)
    _add_textbox(slide, Inches(7.3), y_pos, Inches(4.8), Inches(0.4),
                 f"  {i+1}. {step}", size=13, color=WHITE if i < len(interact_flow)-1 else MAGENTA)
    if i < len(interact_flow) - 1:
        _add_textbox(slide, Inches(7.5), y_pos + Inches(0.38), Inches(4), Inches(0.2),
                     "    ↓", size=11, color=MUTED, align=PP_ALIGN.LEFT)

_add_textbox(slide, Inches(7.1), Inches(6.3), Inches(5), Inches(0.4),
             "交互平面不参与医学风险决策", size=11, color=RED)

# ===========================================================================
# Slide 4: 完整业务流程
# ===========================================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
_set_bg(slide)
_add_title_bar(slide, "完整业务流程", "End-to-End Workflow")

flow_items = [
    ("1", "模式选择", "患者模式 / 临床模式切换", CYAN),
    ("2", "图片上传", "保存本地目录\n不写入 SQLite\n不发送云模型", BLUE),
    ("3", "视觉推理", "DermFM-Zero 编码器\nHAM 监督线性头\n温度缩放 + 拒识", GREEN),
    ("4", "候选输出", "图像质量\nTop-K 候选\n推理耗时", YELLOW),
    ("5", "风险分级", "本地确定性规则\n四级路径生成", AMBER),
    ("6", "症状核验", "结构化追问\n本地 RAG 检索", MAGENTA),
    ("7", "二次校准", "本地规则校准\n只可维持或上调", RED),
    ("8", "对话解释", "ChatBox + MCP\nPrivFlow 出境\n云模型解释", CYAN),
    ("9", "审计存储", "Tool Call\n对话历史\n审计记录", MUTED),
    ("10", "隐私治理", "原始影像不出域\n全链路可追溯", GREEN),
]

# 两行五列
for i, (num, title, desc, color) in enumerate(flow_items):
    col = i % 5
    row = i // 5
    x = Inches(0.6 + col * 2.55)
    y = Inches(1.4 + row * 2.7)
    w = Inches(2.3)
    h = Inches(2.3)

    _add_card(slide, x, y, w, h)
    _add_accent_bar(slide, x, y, w, Inches(0.08), color)
    # 序号圆圈
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.15), y + Inches(0.2), Inches(0.5), Inches(0.5))
    circle.fill.solid()
    circle.fill.fore_color.rgb = color
    circle.line.fill.background()
    circle.shadow.inherit = False
    tf = circle.text_frame
    tf.text = num
    tf.paragraphs[0].font.size = Pt(16)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = BG_DARK
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER

    _add_textbox(slide, x + Inches(0.75), y + Inches(0.25), Inches(1.4), Inches(0.5),
                 title, size=15, color=color, bold=True)
    _add_textbox(slide, x + Inches(0.15), y + Inches(0.85), w - Inches(0.3), Inches(1.3),
                 desc, size=11, color=WHITE)

# ===========================================================================
# Slide 5: 视觉模块
# ===========================================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
_set_bg(slide)
_add_title_bar(slide, "视觉模块", "Frozen Vision Module")

# 左侧：冻结状态
_add_card(slide, Inches(0.85), Inches(1.3), Inches(5.8), Inches(5.5), RGBColor(0x04, 0x14, 0x1E))
_add_accent_bar(slide, Inches(0.85), Inches(1.3), Inches(0.08), Inches(5.5), CYAN)
_add_textbox(slide, Inches(1.1), Inches(1.5), Inches(5), Inches(0.5),
             "DermFM-Zero 冻结主干", size=20, color=CYAN, bold=True)

freeze_items = [
    ("权重：open_clip_pytorch_model.bin", {"size": 13, "color": WHITE}),
    ("大小：约 1.66 GB（永久实体文件）", {"size": 13, "color": WHITE}),
    ("SHA-256 完整性校验", {"size": 13, "color": WHITE}),
    ("独立 Conda 环境：dermfm-zero", {"size": 13, "color": WHITE}),
    ("PyTorch 2.8.0 + CUDA 12.8", {"size": 13, "color": WHITE}),
    ("GPU: RTX 5070 Ti 16GB", {"size": 13, "color": GREEN}),
    ("", {"size": 6, "color": MUTED}),
    ("加载耗时：7,938 ms", {"size": 12, "color": MUTED}),
    ("首次前向：104.60 ms", {"size": 12, "color": MUTED}),
    ("预热前向：18.08 ms", {"size": 12, "color": MUTED}),
    ("峰值显存：1.866 GiB", {"size": 12, "color": MUTED}),
]
_add_multi_text(slide, Inches(1.1), Inches(2.1), Inches(5.3), Inches(4.5), freeze_items, line_spacing=1.3)

# 右侧：HAM7 任务头
_add_card(slide, Inches(6.85), Inches(1.3), Inches(5.8), Inches(5.5), RGBColor(0x1A, 0x04, 0x14))
_add_accent_bar(slide, Inches(6.85), Inches(1.3), Inches(0.08), Inches(5.5), MAGENTA)
_add_textbox(slide, Inches(7.1), Inches(1.5), Inches(5), Inches(0.5),
             "HAM7 监督线性任务头", size=20, color=MAGENTA, bold=True)

ham_items = [
    ("Linear(768, 7)，可训练参数 5,383", {"size": 13, "color": WHITE}),
    ("训练集 8,012 / 验证 500 / 测试 1,503", {"size": 13, "color": WHITE}),
    ("种子 1000，第 124 轮，无类别加权", {"size": 13, "color": WHITE}),
    ("", {"size": 6, "color": MUTED}),
    ("Accuracy: 89.29%", {"size": 16, "color": GREEN, "bold": True}),
    ("Balanced Accuracy: 79.34%", {"size": 16, "color": YELLOW, "bold": True}),
    ("Macro-F1: 80.98%", {"size": 16, "color": CYAN, "bold": True}),
    ("", {"size": 6, "color": MUTED}),
    ("温度缩放 + 低置信度拒识", {"size": 13, "color": WHITE}),
    ("色素痣样例 Top-1: 85.30%", {"size": 13, "color": WHITE}),
    ("皮肤纤维瘤样例仍误判（需保留人工复核）", {"size": 12, "color": RED}),
]
_add_multi_text(slide, Inches(7.1), Inches(2.1), Inches(5.3), Inches(4.5), ham_items, line_spacing=1.3)

# ===========================================================================
# Slide 6: 风险分级与照护路径
# ===========================================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
_set_bg(slide)
_add_title_bar(slide, "风险分级与照护路径", "Triage Levels & Care Pathway")

_add_textbox(slide, Inches(0.85), Inches(1.2), Inches(11), Inches(0.4),
             "系统只允许以下四级路径（本地确定性规则生成）", size=14, color=MUTED)

levels = [
    ("RETAKE", "图像质量不足", "重新拍摄清晰、光线均匀的图片", AMBER),
    ("LOCAL_REVIEW", "基层继续评估", "由基层医生结合病史和现场检查继续评估", CYAN),
    ("REFERRAL", "建议专科转诊", "转诊至具备皮肤专科能力的医疗机构", AMBER),
    ("PRIORITY_REFERRAL", "优先转诊", "优先转诊，接受专科、皮肤镜及病理评估", RED),
]

for i, (level, meaning, action, color) in enumerate(levels):
    y = Inches(1.8 + i * 1.25)
    _add_card(slide, Inches(0.85), y, Inches(11.5), Inches(1.1))
    _add_accent_bar(slide, Inches(0.85), y, Inches(0.1), Inches(1.1), color)

    # 等级编号
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.15), y + Inches(0.25), Inches(0.6), Inches(0.6))
    circle.fill.solid()
    circle.fill.fore_color.rgb = color
    circle.line.fill.background()
    circle.shadow.inherit = False
    tf = circle.text_frame
    tf.text = str(i + 1)
    tf.paragraphs[0].font.size = Pt(20)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = BG_DARK
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER

    _add_textbox(slide, Inches(2.0), y + Inches(0.15), Inches(2.8), Inches(0.4),
                 level, size=16, color=color, bold=True)
    _add_textbox(slide, Inches(2.0), y + Inches(0.55), Inches(2.8), Inches(0.4),
                 meaning, size=13, color=WHITE)
    _add_textbox(slide, Inches(5.2), y + Inches(0.35), Inches(6.8), Inches(0.5),
                 action, size=13, color=MUTED)

_add_textbox(slide, Inches(0.85), Inches(6.8), Inches(11), Inches(0.4),
             "PRIORITY_REFERRAL 表示风险筛查和优先转诊建议，不表示已经确诊癌症。", size=12, color=RED)

# ===========================================================================
# Slide 7: 结构化症状核验 + 二次校准
# ===========================================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
_set_bg(slide)
_add_title_bar(slide, "结构化症状核验与二次校准", "Structured Verification & Calibration")

# 左侧：核验维度
_add_card(slide, Inches(0.85), Inches(1.3), Inches(5.8), Inches(5.5), RGBColor(0x04, 0x14, 0x1E))
_add_accent_bar(slide, Inches(0.85), Inches(1.3), Inches(0.08), Inches(5.5), CYAN)
_add_textbox(slide, Inches(1.1), Inches(1.5), Inches(5), Inches(0.5),
             "核验维度", size=20, color=CYAN, bold=True)

symptoms = [
    "病程和变化速度",
    "疼痛、瘙痒和发热",
    "快速扩散",
    "边缘或颜色变化",
    "感染进展相关表现",
    "呼吸困难、意识异常等系统危险信号",
]
for i, s in enumerate(symptoms):
    y = Inches(2.2 + i * 0.5)
    _add_textbox(slide, Inches(1.3), y, Inches(5), Inches(0.4),
                 f"  ▸  {s}", size=14, color=WHITE)

_add_textbox(slide, Inches(1.1), Inches(5.5), Inches(5), Inches(0.4),
             "云模型只参与白名单问题排序", size=12, color=MUTED)
_add_textbox(slide, Inches(1.1), Inches(5.9), Inches(5), Inches(0.4),
             "最终风险等级由本地规则产生", size=12, color=GREEN)

# 右侧：校准规则
_add_card(slide, Inches(6.85), Inches(1.3), Inches(5.8), Inches(5.5), RGBColor(0x1A, 0x04, 0x14))
_add_accent_bar(slide, Inches(6.85), Inches(1.3), Inches(0.08), Inches(5.5), MAGENTA)
_add_textbox(slide, Inches(7.1), Inches(1.5), Inches(5), Inches(0.5),
             "二次校准规则", size=20, color=MAGENTA, bold=True)

rules = [
    ("校准只能维持或上调风险", {"size": 16, "color": RED, "bold": True}),
    ("不能下调视觉初步风险", {"size": 16, "color": RED, "bold": True}),
    ("", {"size": 8, "color": MUTED}),
    ("R_final = max(R_visual, R_redflag)", {"size": 14, "color": WHITE}),
    ("原因代码生成", {"size": 14, "color": WHITE}),
    ("二次分级由本地规则产生", {"size": 14, "color": WHITE}),
    ("云模型不直接决定结果", {"size": 14, "color": WHITE}),
    ("", {"size": 8, "color": MUTED}),
    ("相关实现：", {"size": 13, "color": MUTED}),
    ("  schemas.py / contextual_service.py", {"size": 12, "color": MUTED}),
    ("  question_bank.json", {"size": 12, "color": MUTED}),
]
_add_multi_text(slide, Inches(7.1), Inches(2.1), Inches(5.3), Inches(4.5), rules, line_spacing=1.3)

# ===========================================================================
# Slide 8: 本地 RAG
# ===========================================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
_set_bg(slide)
_add_title_bar(slide, "本地医学知识检索 (RAG)", "Retrieval-Augmented Generation")

# 检索链
_add_card(slide, Inches(0.85), Inches(1.3), Inches(11.5), Inches(2.2))
_add_textbox(slide, Inches(1.1), Inches(1.45), Inches(11), Inches(0.5),
             "检索链路", size=18, color=CYAN, bold=True)

chain = "患者状态与视觉候选 → 动态 Query → 轻量医学 Dense Retrieval → Top 3~5 → Evidence Assembly → 来源绑定 + 证据不足拒答"
_add_textbox(slide, Inches(1.1), Inches(2.0), Inches(11), Inches(1.3),
             chain, size=14, color=WHITE)

# 知识库
_add_card(slide, Inches(0.85), Inches(3.7), Inches(5.8), Inches(3.2), RGBColor(0x04, 0x14, 0x1E))
_add_accent_bar(slide, Inches(0.85), Inches(3.7), Inches(0.08), Inches(3.2), GREEN)
_add_textbox(slide, Inches(1.1), Inches(3.85), Inches(5), Inches(0.5),
             "知识库 (17 篇文档)", size=18, color=GREEN, bold=True)

docs = [
    "DermFM 能力与数据边界",
    "基层皮肤表现评估",
    "皮损快速变化和危险信号",
    "NICE 可疑皮肤癌转诊",
    "AAD 黑色素瘤 / 角质形成细胞癌",
    "AAD 常见炎症性皮病 / 色素痣",
    "CDC 真菌感染 / 疥疮 / 带状疱疹",
]
for i, d in enumerate(docs):
    _add_textbox(slide, Inches(1.3), Inches(4.4 + i * 0.32), Inches(5), Inches(0.3),
                 f"  ▸  {d}", size=12, color=WHITE)

# 嵌入模型
_add_card(slide, Inches(6.85), Inches(3.7), Inches(5.5), Inches(3.2), RGBColor(0x1A, 0x04, 0x14))
_add_accent_bar(slide, Inches(6.85), Inches(3.7), Inches(0.08), Inches(3.2), MAGENTA)
_add_textbox(slide, Inches(7.1), Inches(3.85), Inches(5), Inches(0.5),
             "嵌入模型", size=18, color=MAGENTA, bold=True)

embed_items = [
    ("IlhaEmbed (本地 checkpoint)", {"size": 14, "color": WHITE}),
    ("model.safetensors，约 152 MB", {"size": 13, "color": WHITE}),
    ("面向中文临床/医学语义", {"size": 13, "color": WHITE}),
    ("", {"size": 8, "color": MUTED}),
    ("证据分类：", {"size": 13, "color": MUTED}),
    ("  风险依据 / 鉴别依据 / 检查依据", {"size": 12, "color": WHITE}),
    ("  转诊依据 / 冲突证据 / 信息缺口", {"size": 12, "color": WHITE}),
    ("", {"size": 8, "color": MUTED}),
    ("来源绑定：文档 / 机构 / 年份 / 章节 / 页码", {"size": 12, "color": MUTED}),
]
_add_multi_text(slide, Inches(7.1), Inches(4.4), Inches(5), Inches(2.5), embed_items, line_spacing=1.2)

# ===========================================================================
# Slide 9: MCP 工具层
# ===========================================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
_set_bg(slide)
_add_title_bar(slide, "MCP 工具层", "Model Context Protocol Tools")

_add_textbox(slide, Inches(0.85), Inches(1.2), Inches(11), Inches(0.4),
             "基于官方 Python MCP SDK，6 个本地工具，与 REST API 共用同一 Service Layer", size=14, color=MUTED)

tools = [
    ("get_visual_evidence", "读取视觉候选、图像质量、Top-K 和拒识状态", CYAN),
    ("get_triage_result", "读取初步分级、原因代码和风险分数", AMBER),
    ("search_guideline", "检索本地 RAG 指南和证据", GREEN),
    ("get_care_pathway", "读取照护路径和建议下一步", BLUE),
    ("get_privacy_report", "读取隐私出境和审计状态", MAGENTA),
    ("get_contextual_assessment", "读取症状二次校准结果", YELLOW),
]

for i, (name, desc, color) in enumerate(tools):
    col = i % 2
    row = i // 2
    x = Inches(0.85 + col * 6.2)
    y = Inches(1.8 + row * 1.6)
    w = Inches(5.8)
    h = Inches(1.4)

    _add_card(slide, x, y, w, h)
    _add_accent_bar(slide, x, y, Inches(0.08), h, color)

    # 工具编号
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.2), y + Inches(0.35), Inches(0.7), Inches(0.7))
    circle.fill.solid()
    circle.fill.fore_color.rgb = color
    circle.line.fill.background()
    circle.shadow.inherit = False
    tf = circle.text_frame
    tf.text = str(i + 1)
    tf.paragraphs[0].font.size = Pt(22)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = BG_DARK
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER

    _add_textbox(slide, x + Inches(1.1), y + Inches(0.2), Inches(4.5), Inches(0.5),
                 name, size=16, color=color, bold=True, font_name="Consolas")
    _add_textbox(slide, x + Inches(1.1), y + Inches(0.7), Inches(4.5), Inches(0.6),
                 desc, size=12, color=WHITE)

_add_textbox(slide, Inches(0.85), Inches(6.6), Inches(11), Inches(0.5),
             "MCP 用于标准化 Agent 对本地能力的调用，不提高医疗准确率", size=12, color=MUTED)

# ===========================================================================
# Slide 10: PrivFlow 隐私治理
# ===========================================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
_set_bg(slide)
_add_title_bar(slide, "PrivFlow 隐私治理", "Privacy Egress Governance")

# 禁止外发
_add_card(slide, Inches(0.85), Inches(1.3), Inches(5.8), Inches(5.5), RGBColor(0x1A, 0x04, 0x10))
_add_accent_bar(slide, Inches(0.85), Inches(1.3), Inches(0.08), Inches(5.5), RED)
_add_textbox(slide, Inches(1.1), Inches(1.5), Inches(5), Inches(0.5),
             "禁止外发", size=22, color=RED, bold=True)

forbidden = [
    "原始医疗影像和图像字节",
    "Base64 图片或 data:image/ 内容",
    "本地图片路径",
    "EXIF 等原始文件信息",
    "姓名、电话、身份证号和地址",
    "患者身份映射信息",
]
for i, item in enumerate(forbidden):
    _add_textbox(slide, Inches(1.3), Inches(2.2 + i * 0.55), Inches(5), Inches(0.4),
                 f"  ✗  {item}", size=15, color=RED)

_add_textbox(slide, Inches(1.1), Inches(5.8), Inches(5), Inches(0.4),
             "所有拦截动作写入审计记录", size=12, color=MUTED)

# 允许外发
_add_card(slide, Inches(6.85), Inches(1.3), Inches(5.8), Inches(5.5), RGBColor(0x04, 0x14, 0x10))
_add_accent_bar(slide, Inches(6.85), Inches(1.3), Inches(0.08), Inches(5.5), GREEN)
_add_textbox(slide, Inches(7.1), Inches(1.5), Inches(5), Inches(0.5),
             "允许外发（检查后）", size=22, color=GREEN, bold=True)

allowed = [
    "匿名病例标识",
    "脱敏后的结构化视觉结果",
    "分级与照护路径",
    "结构化症状核验结果",
    "本地 RAG 摘要和指南来源",
    "用户当前问题",
]
for i, item in enumerate(allowed):
    _add_textbox(slide, Inches(7.3), Inches(2.2 + i * 0.55), Inches(5), Inches(0.4),
                 f"  ✓  {item}", size=15, color=GREEN)

_add_textbox(slide, Inches(7.1), Inches(5.8), Inches(5), Inches(0.4),
             "允许、拦截和清洗动作写入审计", size=12, color=MUTED)

# ===========================================================================
# Slide 11: 对话 Agent 与云模型
# ===========================================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
_set_bg(slide)
_add_title_bar(slide, "对话 Agent 与云模型", "Chat Agent & Cloud LLM")

_add_textbox(slide, Inches(0.85), Inches(1.2), Inches(11), Inches(0.4),
             "DermaGuardChatAgent — 有限步检索 Agent，不使用多 Agent / ReAct 无限循环 / 复杂 Planner",
             size=14, color=MUTED)

# 能力列表
_add_card(slide, Inches(0.85), Inches(1.8), Inches(5.8), Inches(5.0))
_add_accent_bar(slide, Inches(0.85), Inches(1.8), Inches(0.08), Inches(5.0), CYAN)
_add_textbox(slide, Inches(1.1), Inches(2.0), Inches(5), Inches(0.5),
             "Agent 能力", size=20, color=CYAN, bold=True)

abilities = [
    "解释视觉候选和分级结果",
    "解释转诊原因和下一步路径",
    "回答隐私问题",
    "读取症状二次校准结果",
    "查询本地指南",
    "保留最近 12 条对话上下文",
    "云模型失败时降级到本地 Mock",
]
for i, a in enumerate(abilities):
    _add_textbox(slide, Inches(1.3), Inches(2.7 + i * 0.5), Inches(5), Inches(0.4),
                 f"  ▸  {a}", size=14, color=WHITE)

# 模型目录
_add_card(slide, Inches(6.85), Inches(1.8), Inches(5.8), Inches(5.0))
_add_accent_bar(slide, Inches(6.85), Inches(1.8), Inches(0.08), Inches(5.0), MAGENTA)
_add_textbox(slide, Inches(7.1), Inches(2.0), Inches(5), Inches(0.5),
             "模型目录", size=20, color=MAGENTA, bold=True)

models = [
    ("智谱 GLM-5.3-Flash", "当前主力", GREEN),
    ("智谱 GLM-4-Flash", "备用", CYAN),
    ("Gemini Flash", "需单独配置密钥", YELLOW),
    ("Groq-Llama3", "已停供，不可用", RED),
]
for i, (name, status, color) in enumerate(models):
    y = Inches(2.7 + i * 0.55)
    _add_textbox(slide, Inches(7.3), y, Inches(3.5), Inches(0.4),
                 f"  ▸  {name}", size=15, color=color, bold=True)
    _add_textbox(slide, Inches(10.5), y, Inches(2), Inches(0.4),
                 status, size=12, color=MUTED)

_add_textbox(slide, Inches(7.1), Inches(5.2), Inches(5), Inches(0.4),
             "密钥保存在本地 .env，不写入源码", size=12, color=RED)
_add_textbox(slide, Inches(7.1), Inches(5.6), Inches(5), Inches(0.4),
             "支持简明回答 / 详细归因双档切换", size=12, color=MUTED)
_add_textbox(slide, Inches(7.1), Inches(6.0), Inches(5), Inches(0.4),
             "逐字显示（非原生流式）", size=12, color=MUTED)

# ===========================================================================
# Slide 12: REST API
# ===========================================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
_set_bg(slide)
_add_title_bar(slide, "REST API", "FastAPI Endpoints")

_add_textbox(slide, Inches(0.85), Inches(1.2), Inches(11), Inches(0.4),
             "后端使用 FastAPI，主要接口如下", size=14, color=MUTED)

headers = ["方法", "路径", "作用"]
rows = [
    ["GET", "/health", "健康检查"],
    ["POST", "/v1/cases", "创建病例并保存本地图片"],
    ["POST", "/v1/cases/{id}/analyze", "执行视觉、分级、RAG 和照护路径"],
    ["GET", "/v1/cases/{id}/visual", "读取视觉证据"],
    ["GET", "/v1/cases/{id}/triage", "读取初步分级"],
    ["GET", "/v1/cases/{id}/care-pathway", "读取照护路径"],
    ["GET", "/v1/cases/{id}/privacy", "读取隐私报告"],
    ["GET", "/v1/cases/{id}/cloud-payload", "查看允许发送到云端的脱敏载荷"],
    ["POST", "/v1/cases/{id}/contextual-intake", "生成结构化追问计划"],
    ["POST", "/v1/cases/{id}/contextual-calibration", "完成症状二次校准"],
    ["POST", "/v1/chat", "调用对话 Agent"],
    ["GET", "/v1/audit", "读取审计记录"],
]
_add_table(slide, Inches(0.85), Inches(1.8), Inches(11.5), Inches(5.2), headers, rows)

# ===========================================================================
# Slide 13: 数据存储与项目目录
# ===========================================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
_set_bg(slide)
_add_title_bar(slide, "数据存储与项目目录", "Storage & Project Structure")

# 数据存储
_add_card(slide, Inches(0.85), Inches(1.3), Inches(5.8), Inches(5.5))
_add_accent_bar(slide, Inches(0.85), Inches(1.3), Inches(0.08), Inches(5.5), CYAN)
_add_textbox(slide, Inches(1.1), Inches(1.5), Inches(5), Inches(0.5),
             "数据存储 (SQLite)", size=20, color=CYAN, bold=True)

_add_textbox(slide, Inches(1.1), Inches(2.1), Inches(5), Inches(0.4),
             "data/dermaguard.db", size=14, color=GREEN, font_name="Consolas")

db_items = [
    "匿名病例记录",
    "图片本地路径和 SHA-256",
    "视觉结果",
    "分级结果",
    "Care Pathway",
    "症状二次校准结果",
    "脱敏对话历史",
    "MCP Tool Call",
    "PrivFlow 审计事件",
]
for i, item in enumerate(db_items):
    _add_textbox(slide, Inches(1.3), Inches(2.6 + i * 0.4), Inches(5), Inches(0.35),
                 f"  ▸  {item}", size=12, color=WHITE)

_add_textbox(slide, Inches(1.1), Inches(6.3), Inches(5), Inches(0.4),
             "数据库不保存图片 BLOB", size=12, color=RED)

# 项目目录
_add_card(slide, Inches(6.85), Inches(1.3), Inches(5.8), Inches(5.5))
_add_accent_bar(slide, Inches(6.85), Inches(1.3), Inches(0.08), Inches(5.5), MAGENTA)
_add_textbox(slide, Inches(7.1), Inches(1.5), Inches(5), Inches(0.5),
             "项目目录", size=20, color=MAGENTA, bold=True)

dirs = [
    ("backend/app/api/", "FastAPI 路由"),
    ("backend/app/care/", "四级 Care Pathway"),
    ("backend/app/chat/", "Chat Agent 与 Provider"),
    ("backend/app/contextual/", "症状核验"),
    ("backend/app/knowledge/", "规则 / 问题库 / RAG"),
    ("backend/app/mcp/", "MCP Server / Tools"),
    ("backend/app/services/", "REST/MCP 共用业务"),
    ("data/cases/", "本地病例图片"),
    ("ui/", "Streamlit 页面 + ChatBox"),
    ("weights/", "本地实体权重"),
]
for i, (path, desc) in enumerate(dirs):
    y = Inches(2.1 + i * 0.45)
    _add_textbox(slide, Inches(7.3), y, Inches(2.8), Inches(0.35),
                 path, size=11, color=GREEN, font_name="Consolas")
    _add_textbox(slide, Inches(10.2), y, Inches(2.2), Inches(0.35),
                 desc, size=11, color=MUTED)

# ===========================================================================
# Slide 14: 测试与验证
# ===========================================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
_set_bg(slide)
_add_title_bar(slide, "测试与验证状态", "Testing & Validation")

_add_textbox(slide, Inches(0.85), Inches(1.2), Inches(11), Inches(0.4),
             "截至 2026-10-11 本机验证结果", size=14, color=MUTED)

# 测试结果卡片
test_cards = [
    ("后端测试", "63 项通过", GREEN),
    ("视觉 Provider", "7 项通过", GREEN),
    ("HAM 任务头前向", "色素痣 Top-1 85.30%", CYAN),
    ("RAG 检索链", "本地实体权重验证", GREEN),
    ("FastAPI 8000", "健康检查通过", CYAN),
    ("Streamlit 8501", "健康检查通过", CYAN),
]

for i, (title, result, color) in enumerate(test_cards):
    col = i % 3
    row = i // 3
    x = Inches(0.85 + col * 4.0)
    y = Inches(1.8 + row * 2.0)
    w = Inches(3.7)
    h = Inches(1.7)

    _add_card(slide, x, y, w, h)
    _add_accent_bar(slide, x, y, w, Inches(0.08), color)
    _add_textbox(slide, x + Inches(0.2), y + Inches(0.3), w - Inches(0.3), Inches(0.5),
                 title, size=16, color=MUTED, bold=True)
    _add_textbox(slide, x + Inches(0.2), y + Inches(0.8), w - Inches(0.3), Inches(0.7),
                 result, size=20, color=color, bold=True)

_add_textbox(slide, Inches(0.85), Inches(6.0), Inches(11), Inches(0.5),
             "后端测试注入本地 Mock Provider，不消耗云端额度", size=13, color=AMBER)
_add_textbox(slide, Inches(0.85), Inches(6.4), Inches(11), Inches(0.5),
             "第三方依赖存在弃用警告，不影响当前前向测试", size=13, color=MUTED)

# ===========================================================================
# Slide 15: 完成度与待完善
# ===========================================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
_set_bg(slide)
_add_title_bar(slide, "完成度与后续规划", "Progress & Roadmap")

# 完成度
_add_card(slide, Inches(0.85), Inches(1.3), Inches(5.8), Inches(5.5), RGBColor(0x04, 0x14, 0x1E))
_add_accent_bar(slide, Inches(0.85), Inches(1.3), Inches(0.08), Inches(5.5), GREEN)
_add_textbox(slide, Inches(1.1), Inches(1.5), Inches(5), Inches(0.5),
             "已完成（约 92%）", size=20, color=GREEN, bold=True)

done_items = [
    "冻结视觉模型真实权重和单图前向",
    "HAM 监督任务头接入网页推理链",
    "温度缩放和低置信度拒识",
    "四级本地分级与 Care Pathway",
    "结构化症状追问和二次校准",
    "轻量医学 Dense Retrieval + 证据组装",
    "6 个 MCP 工具和真实 MCP Server",
    "PrivFlow 出境检查与审计",
    "智谱云模型 + 本地 Mock 降级",
    "悬浮 ChatBox + Markdown 渲染",
    "SQLite 全链路存储",
    "后端及视觉 Provider 自动化测试",
]
for i, item in enumerate(done_items):
    col = i % 2
    row = i // 2
    x = Inches(1.1 + col * 2.8)
    y = Inches(2.1 + row * 0.4)
    _add_textbox(slide, x, y, Inches(2.7), Inches(0.35),
                 f"  ✓  {item}", size=10, color=WHITE)

# 待完善
_add_card(slide, Inches(6.85), Inches(1.3), Inches(5.8), Inches(5.5), RGBColor(0x1A, 0x10, 0x04))
_add_accent_bar(slide, Inches(6.85), Inches(1.3), Inches(0.08), Inches(5.5), AMBER)
_add_textbox(slide, Inches(7.1), Inches(1.5), Inches(5), Inches(0.5),
             "仍需完善", size=20, color=AMBER, bold=True)

todo_items = [
    "1. 扩展皮肤科权威医学资料",
    "   补齐机构、年份、版本、章节元数据",
    "2. 建立真实用户问法测试集",
    "   评估意图路由和归因完整度",
    "3. 升级为服务端原生流式传输",
    "   替代逐字显示",
    "4. 增加浏览器交互测试",
    "   持续运行的云 API 集成测试",
    "5. 生产级补充",
    "   临床验证 / 鉴权 / 加密 / 监控",
    "   备份恢复 / 并发 / 合规评审",
]
_add_multi_text(slide, Inches(7.1), Inches(2.1), Inches(5.3), Inches(4.5),
                [(t, {"size": 13, "color": WHITE}) for t in todo_items], line_spacing=1.2)

# ===========================================================================
# Slide 16: 安全边界
# ===========================================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
_set_bg(slide)
_add_title_bar(slide, "安全边界", "Safety Boundaries")

_add_textbox(slide, Inches(0.85), Inches(1.2), Inches(11), Inches(0.4),
             "任何页面、对话或报告不得出现以下行为", size=14, color=MUTED)

boundaries = [
    ("声称患者已经确诊某种皮肤病或癌症", RED),
    ("声称患者确定没有疾病或无需就医", RED),
    ("把视觉候选分数解释为患病概率", RED),
    ("由云模型修改本地分级或视觉结果", RED),
    ("自动开药、给出处方或替代医生制定治疗方案", RED),
    ("把原始图片、图片编码、本地路径或患者身份发送给云模型", RED),
]

for i, (text, color) in enumerate(boundaries):
    y = Inches(1.8 + i * 0.7)
    _add_card(slide, Inches(0.85), y, Inches(11.5), Inches(0.6))
    _add_accent_bar(slide, Inches(0.85), y, Inches(0.1), Inches(0.6), color)
    _add_textbox(slide, Inches(1.2), y + Inches(0.12), Inches(10.5), Inches(0.4),
                 f"  ✗  {text}", size=15, color=color, bold=True)

# 危险信号提示
_add_card(slide, Inches(0.85), Inches(6.3), Inches(11.5), Inches(0.9), RGBColor(0x1A, 0x04, 0x10))
_add_accent_bar(slide, Inches(0.85), Inches(6.3), Inches(0.1), Inches(0.9), AMBER)
_add_textbox(slide, Inches(1.2), Inches(6.45), Inches(10.5), Inches(0.6),
             "⚠  出现呼吸困难、意识异常、严重出血、严重过敏等危险信号时，应立即寻求线下急诊",
             size=14, color=AMBER, bold=True)

# ===========================================================================
# Slide 17: 启动方式
# ===========================================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
_set_bg(slide)
_add_title_bar(slide, "启动方式", "How to Run")

# 两个环境
_add_card(slide, Inches(0.85), Inches(1.3), Inches(11.5), Inches(1.5))
_add_multi_text(slide, Inches(1.1), Inches(1.45), Inches(11), Inches(1.3), [
    ("两个 Conda 环境：", {"size": 16, "color": CYAN, "bold": True}),
    ("  • dermaguard — FastAPI、Streamlit、MCP、RAG 和 Chat Agent", {"size": 14, "color": WHITE}),
    ("  • dermfm-zero — 冻结 DermFM-Zero 视觉推理", {"size": 14, "color": WHITE}),
])

# 一键启动
_add_card(slide, Inches(0.85), Inches(3.1), Inches(11.5), Inches(1.5), RGBColor(0x04, 0x14, 0x1E))
_add_accent_bar(slide, Inches(0.85), Inches(3.1), Inches(0.08), Inches(1.5), GREEN)
_add_textbox(slide, Inches(1.1), Inches(3.25), Inches(5), Inches(0.4),
             "一键启动", size=16, color=GREEN, bold=True)
_add_textbox(slide, Inches(1.1), Inches(3.7), Inches(11), Inches(0.5),
             "bash scripts/run_mock_demo.sh", size=16, color=GREEN, font_name="Consolas")

# 访问地址
_add_card(slide, Inches(0.85), Inches(4.9), Inches(11.5), Inches(1.5), RGBColor(0x1A, 0x04, 0x14))
_add_accent_bar(slide, Inches(0.85), Inches(4.9), Inches(0.08), Inches(1.5), MAGENTA)
_add_textbox(slide, Inches(1.1), Inches(5.05), Inches(5), Inches(0.4),
             "访问地址", size=16, color=MAGENTA, bold=True)
_add_multi_text(slide, Inches(1.1), Inches(5.5), Inches(11), Inches(0.8), [
    ("Web 页面：http://127.0.0.1:8501", {"size": 15, "color": WHITE, "font_name": "Consolas"}),
    ("API 文档：http://127.0.0.1:8000/docs", {"size": 15, "color": WHITE, "font_name": "Consolas"}),
])

_add_textbox(slide, Inches(0.85), Inches(6.7), Inches(11), Inches(0.4),
             "MCP stdio Server 单独启动：bash scripts/run_mcp_server.sh", size=12, color=MUTED)

# ===========================================================================
# Slide 18: 结尾
# ===========================================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
_set_bg(slide)

_add_accent_bar(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.06), CYAN)
_add_accent_bar(slide, Inches(0), Inches(7.44), Inches(13.333), Inches(0.06), MAGENTA)

_add_textbox(slide, Inches(1.5), Inches(2.2), Inches(10), Inches(1.0),
             "DermaGuard", size=48, color=CYAN, bold=True, align=PP_ALIGN.CENTER)
_add_textbox(slide, Inches(1.5), Inches(3.3), Inches(10), Inches(0.6),
             "基层皮肤病灶智能筛查与分级转诊", size=22, color=WHITE, align=PP_ALIGN.CENTER)

_add_textbox(slide, Inches(1.5), Inches(4.5), Inches(10), Inches(0.5),
             "Local-First · Privacy-Preserving · AI-Assisted · MCP-Powered",
             size=16, color=MAGENTA, align=PP_ALIGN.CENTER)

_add_textbox(slide, Inches(1.5), Inches(5.5), Inches(10), Inches(0.5),
             "本系统仅用于辅助筛查和分级转诊，不构成医学诊断。",
             size=14, color=MUTED, align=PP_ALIGN.CENTER)
_add_textbox(slide, Inches(1.5), Inches(6.2), Inches(10), Inches(0.5),
             "Thank You", size=28, color=CYAN, bold=True, align=PP_ALIGN.CENTER)

# ---------------------------------------------------------------------------
# 保存
# ---------------------------------------------------------------------------
output_path = "/media/lu/CE387FD6387FBC51/科研项目/医疗agent/DermaGuard_项目介绍.pptx"
prs.save(output_path)
print(f"PPT 已生成：{output_path}")
print(f"共 {len(prs.slides)} 页幻灯片")
